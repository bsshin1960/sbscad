import sys
import subprocess

def install(package):
    subprocess.check_call([sys.executable, "-m", "pip", "install", package])

try:
    import docx
except ImportError:
    install('python-docx')
    import docx

doc = docx.Document()
doc.add_heading('3D CAD 모델링 제작 프로그램 개발 상세 계획서 (CATIA 스타일)', 0)

doc.add_paragraph('이 계획서는 단순한 3D 뷰어가 아닌, 사용자가 직접 2D 스케치를 작성하고 이를 바탕으로 3D 입체 형상을 제작(Authoring)할 수 있는 파라메트릭 3D CAD 모델링 제작 프로그램을 개발하기 위한 상세 계획입니다.')

doc.add_heading('1. 개발 목적 및 핵심 워크플로우', level=2)
doc.add_paragraph('본 프로그램의 핵심은 산업용 CAD(CATIA, SolidWorks 등)의 표준 방식인 "2D 스케치 기반 3D 피처(Feature) 생성"입니다.')
doc.add_paragraph('Workflow: 2D 평면 선택 -> 2D 스케치 그리기(직선, 곡선, 원, 사각형) -> 3D 피처로 변환(돌출, 회전) -> 형상 다듬기(챔퍼, 라운드) -> 작업 트리에서 히스토리 관리', style='List Bullet')

doc.add_heading('2. 개발 언어 및 기술 스택 선정', level=2)
doc.add_paragraph('주요 프로그래밍 언어: Python (파이썬)', style='List Bullet')
doc.add_paragraph('UI 프레임워크: PyQt6\n선정 이유: CATIA와 유사한 트리 뷰, 속성창, 스케치 툴바 등을 정교하게 배치할 수 있습니다.', style='List Bullet')
doc.add_paragraph('3D 커널 및 렌더링 엔진: PyVista(VTK) / pythonOCC(OpenCASCADE)\n선정 이유: 2D 스케치 데이터를 수학적으로 계산하여 3D 솔리드로 변환하고(Extrude, Revolve, Fillet 등) 렌더링하기 위한 필수 엔진입니다.', style='List Bullet')

doc.add_heading('3. 주요 개발 기능 및 조작법 상세 (1차 버전)', level=2)
doc.add_heading('A. 2D 스케처 (Sketcher) 모드', level=3)
doc.add_paragraph('사용자가 3D 공간 상의 특정 평면(XY, YZ, ZX)을 선택하면, 2D 도면을 그릴 수 있는 스케치 모드로 진입합니다.')
doc.add_paragraph('직선 (Line): 두 점을 지정하여 직선 선분을 생성합니다.', style='List Bullet')
doc.add_paragraph('곡선 (Curve/Spline): 여러 제어점을 클릭하여 부드러운 자유 곡선을 생성합니다.', style='List Bullet')
doc.add_paragraph('원 (Circle): 중심점과 반지름을 지정하여 원을 생성합니다.', style='List Bullet')
doc.add_paragraph('사각형 (Rectangle): 두 대각선 꼭짓점을 지정하여 사각형을 생성합니다.', style='List Bullet')

doc.add_heading('B. 3D 모델링 제작 (3D Features)', level=3)
doc.add_paragraph('완성된 2D 스케치의 닫힌 단면(Profile)을 활용하여 3D 형상을 만들어내거나, 기존 3D 입체를 변형합니다.')
doc.add_paragraph('기초 형상 생성 피처 (Sketch-Based Features):', style='List Bullet')
doc.add_paragraph('  - 돌출 (Pad / Extrude): 2D 스케치를 수직 방향으로 밀어내어(길이 지정) 3D 입체를 생성합니다.')
doc.add_paragraph('  - 회전 (Shaft / Revolve): 2D 스케치를 중심축을 기준으로 360도 회전시켜 원통형이나 원뿔형 입체를 생성합니다.')
doc.add_paragraph('형상 다듬기 피처 (Dress-up Features):', style='List Bullet')
doc.add_paragraph('  - 라운드/모깎기 (Round / Fillet): 3D 모델의 날카로운 모서리(Edge)나 꼭짓점을 선택하여, 지정한 반지름(Radius)만큼 둥글게 깎아내거나 살을 붙입니다.')
doc.add_paragraph('  - 챔퍼/모따기 (Chamfer): 3D 모델의 모서리를 선택하여, 지정한 거리(Distance)나 각도(Angle)로 비스듬하게 평면으로 잘라냅니다(경사면 생성).')

doc.add_heading('C. 3D 뷰어 마우스 조작법 (View Control)', level=3)
doc.add_paragraph('모델링 과정을 직관적으로 확인할 수 있는 마우스 조작법을 제공합니다.')
doc.add_paragraph('화면 회전 (Rotate): 마우스 왼쪽 버튼 누른 채 드래그', style='List Bullet')
doc.add_paragraph('화면 이동 (Pan): 마우스 휠(가운데) 버튼 누른 채 드래그 (또는 Shift + 왼쪽 버튼 드래그)', style='List Bullet')
doc.add_paragraph('확대/축소 (Zoom): 마우스 휠 스크롤 (또는 마우스 오른쪽 버튼 누른 채 드래그 상하 이동)', style='List Bullet')
doc.add_paragraph('화면 초기화 (Fit All In): 모델이 화면 중앙에 꽉 차게 재정렬', style='List Bullet')

doc.add_heading('D. 작업 트리 (Specification Tree) 및 히스토리 관리', level=3)
doc.add_paragraph('좌측 패널에 생성된 \'기준 평면 -> 2D 스케치 -> 3D 피처(돌출/회전) -> 다듬기(라운드/챔퍼)\'의 작업 이력이 트리 계층 구조로 표시됩니다.', style='List Bullet')
doc.add_paragraph('이를 통해 언제든지 스케치를 재수정하거나 돌출 길이, 모깎기 반지름 등을 다시 편집(Parametric Editing)할 수 있습니다.', style='List Bullet')

doc.add_heading('4. 전체 메뉴 및 UI 구조 설계', level=2)
doc.add_paragraph('메뉴바: File(새로만들기, 저장), Insert(Sketcher, Pad, Shaft, Fillet, Chamfer)', style='List Bullet')
doc.add_paragraph('좌측 패널: 작업 히스토리를 보여주는 트리 창 (Specification Tree)', style='List Bullet')
doc.add_paragraph('중앙 뷰포트: 2D 스케치 및 3D 모델링 메인 작업 공간 (CATIA 특유의 짙은 그라데이션 배경)', style='List Bullet')
doc.add_paragraph('툴바: 스케치 도구 및 3D 변환, 모깎기/모따기 아이콘 배치', style='List Bullet')

try:
    doc.save('c:\\Temp\\Antigrvity\\sbscad\\Implementation Plan.docx')
    print("Docx created successfully.")
except PermissionError:
    doc.save('c:\\Temp\\Antigrvity\\sbscad\\Implementation Plan_Updated.docx')
    print("Docx created successfully as 'Implementation Plan_Updated.docx' because the original was open.")
